from __future__ import annotations
from pathlib import Path
import concurrent.futures
import mimetypes


class AttachmentParserError(Exception):
    pass


class BaseParser:
    MIME_TYPES: set[str] = set()
    def parse(self, file_path: str) -> str:
        raise NotImplementedError


class TxtParser(BaseParser):
    MIME_TYPES = {'text/plain'}
    def parse(self, file_path: str) -> str:
        raw = Path(file_path).read_bytes()
        for enc in ('utf-8', 'latin-1', 'utf-16'):
            try:
                return raw.decode(enc)
            except Exception:
                continue
        return raw.decode('utf-8', errors='ignore')


class PdfParser(BaseParser):
    MIME_TYPES = {'application/pdf'}
    def parse(self, file_path: str) -> str:
        try:
            from pypdf import PdfReader
        except Exception as e:
            raise AttachmentParserError('pdf_parser_dependency_missing') from e
        reader = PdfReader(file_path, strict=False)
        return '\n'.join([(p.extract_text() or '') for p in reader.pages]).strip()


class DocxParser(BaseParser):
    MIME_TYPES = {'application/vnd.openxmlformats-officedocument.wordprocessingml.document'}
    def parse(self, file_path: str) -> str:
        try:
            import docx
        except Exception as e:
            raise AttachmentParserError('docx_parser_dependency_missing') from e
        d = docx.Document(file_path)
        return '\n'.join([p.text for p in d.paragraphs]).strip()


class XlsxParser(BaseParser):
    MIME_TYPES = {'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'}
    def parse(self, file_path: str) -> str:
        try:
            from openpyxl import load_workbook
        except Exception as e:
            raise AttachmentParserError('xlsx_parser_dependency_missing') from e
        wb = load_workbook(file_path, data_only=True, read_only=True)
        lines = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    lines.append(' | '.join(cells))
        return '\n'.join(lines).strip()


class PptxParser(BaseParser):
    MIME_TYPES = {'application/vnd.openxmlformats-officedocument.presentationml.presentation'}
    def parse(self, file_path: str) -> str:
        try:
            from pptx import Presentation
        except Exception as e:
            raise AttachmentParserError('pptx_parser_dependency_missing') from e
        prs = Presentation(file_path)
        lines = []
        for s in prs.slides:
            for shape in s.shapes:
                text = getattr(shape, 'text', None)
                if text:
                    lines.append(text)
        return '\n'.join(lines).strip()


PARSER_REGISTRY: list[BaseParser] = [TxtParser(), PdfParser(), DocxParser(), XlsxParser(), PptxParser()]


def _run_parser(parser: BaseParser, file_path: str) -> str:
    return parser.parse(file_path)


def parse_attachment(file_path: str, mime_type: str, timeout_seconds: int = 10, max_chars: int = 250_000) -> str:
    guessed, _ = mimetypes.guess_type(file_path)
    if guessed and mime_type != guessed and mime_type != 'application/octet-stream':
        raise AttachmentParserError('mime_validation_failed')

    parser = next((p for p in PARSER_REGISTRY if mime_type in p.MIME_TYPES), None)
    if parser is None:
        raise AttachmentParserError('unsupported_parser_mime')

    try:
        with concurrent.futures.ThreadPoolExecutor(max_workers=1) as ex:
            fut = ex.submit(_run_parser, parser, file_path)
            text = fut.result(timeout=timeout_seconds)
    except concurrent.futures.TimeoutError as e:
        raise AttachmentParserError('parser_timeout') from e
    except Exception as e:
        raise AttachmentParserError(f'parser_failed:{type(e).__name__}') from e

    if len(text) > max_chars:
        text = text[:max_chars]
    return text
